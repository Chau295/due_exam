# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import random
import base64
import io
import os
import logging
from base64 import b64encode
from typing import List, Dict, Any
from uuid import uuid4

logger = logging.getLogger(__name__)

from django import forms
from django.contrib import messages
from django.contrib.auth import get_user_model, login
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.password_validation import validate_password
from django.core.cache import cache
from django.core.exceptions import PermissionDenied, ValidationError
from django.db.models import Sum, Avg, Count
from django.http import (
    JsonResponse,
    HttpRequest,
    HttpResponse,
    HttpResponseBadRequest,
)
from django.shortcuts import get_object_or_404, render, redirect
from django.templatetags.static import static
from django.utils import timezone
from django.utils.safestring import mark_safe
from django.views.decorators.http import require_POST, require_GET
from django.conf import settings

from asgiref.sync import async_to_sync
from channels.layers import get_channel_layer


from .forms import (
    LecturerManualQuestionForm,
    LecturerMaterialUploadForm,
)
from .models import (
    Subject,
    Question,
    ExamSession,
    ExamResult,
    SupplementaryResult,
    UserProfile,
    ExamCode,
    LectureMaterial,
    ExamRoom,
    ExamSessionGroup,
    ExamSessionRoom,
    DifficultyLevel,
)

from django.contrib.auth.decorators import login_required
from django.shortcuts import get_object_or_404
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn



User = get_user_model()

# ===========================
# GIẢNG VIÊN - NEW INDEPENDENT SCREENS
# ===========================

@login_required
def post_login_redirect(request):
    if not hasattr(request.user, 'userprofile'):
        UserProfile.objects.create(user=request.user)

    if request.user.userprofile.is_lecturer:
        return redirect('qna:lecturer_dashboard')
    else:
        return redirect('qna:dashboard')


# ===========================
# 1. DASHBOARD - System Overview
# ===========================

@login_required
def lecturer_dashboard(request):
    if not request.user.userprofile.is_lecturer:
        return redirect('qna:dashboard')

    subjects = request.user.userprofile.subjects_taught.all()
    selected_subject_id = request.GET.get('subject_id')

    total_subjects = subjects.count()
    total_sessions = ExamSession.objects.filter(subject__in=subjects).count()

    if selected_subject_id:
        try:
            selected_subject = subjects.get(id=selected_subject_id)
            sessions = ExamSession.objects.filter(
                subject=selected_subject
            ).select_related('subject', 'user', 'user__userprofile').order_by('-created_at')[:20]
        except (Subject.DoesNotExist, ValueError):
            selected_subject = None
            sessions = []
    else:
        selected_subject = None
        sessions = ExamSession.objects.filter(
            subject__in=subjects
        ).select_related('subject', 'user', 'user__userprofile').order_by('-created_at')[:20]

    context = {
        'subjects': subjects,
        'selected_subject': selected_subject,
        'total_subjects': total_subjects,
        'total_sessions': total_sessions,
        'recent_sessions': sessions,
    }
    return render(request, 'qna/lecturer/lecturer_dashboard.html', context)


@login_required
def lecturer_subject_list(request):
    if not request.user.userprofile.is_lecturer:
        return redirect('qna:dashboard')

    subjects = request.user.userprofile.subjects_taught.all().annotate(
        exam_count=Count('examsession', distinct=True)
    ).order_by('name')

    return render(request, 'qna/lecturer/lecturer_subject_list.html', {'subjects': subjects})


@login_required
def lecturer_subject_workspace(request, subject_code):
    if not request.user.userprofile.is_lecturer:
        return redirect('qna:dashboard')

    subject = get_object_or_404(
        request.user.userprofile.subjects_taught,
        subject_code=subject_code
    )

    sessions_in_subject = ExamSession.objects.filter(subject=subject)
    total_exams = sessions_in_subject.count()
    completed_exams = sessions_in_subject.filter(is_completed=True).count()

    upcoming_groups = ExamSessionGroup.objects.filter(
        subject=subject,
        exam_date__gte=timezone.now()
    ).order_by('exam_date')[:5]

    approved_codes = ExamCode.objects.filter(subject=subject, is_approved=True).count()

    context = {
        'subject': subject,
        'total_exams': total_exams,
        'completed_exams': completed_exams,
        'upcoming_groups': upcoming_groups,
        'approved_codes': approved_codes,
    }
    return render(request, 'qna/lecturer/lecturer_subject_workspace.html', context)


@login_required
def lecturer_exam_codes_screen(request):
    if not request.user.userprofile.is_lecturer:
        return redirect('qna:dashboard')

    subjects = request.user.userprofile.subjects_taught.all()
    selected_subject_id = request.GET.get('subject_id')

    if not selected_subject_id:
        selected_subject = subjects.first()
    else:
        try:
            selected_subject = subjects.get(id=selected_subject_id)
        except (Subject.DoesNotExist, ValueError):
            selected_subject = subjects.first()

    all_questions = Question.objects.filter(subject=selected_subject).order_by('question_id_in_barem')
    materials = LectureMaterial.objects.filter(subject=selected_subject).order_by('-uploaded_at')

    exam_codes = ExamCode.objects.filter(subject=selected_subject).select_related(
        'question_easy', 'question_medium', 'question_hard'
    ).order_by('-created_at')

    pending_exam_codes = ExamCode.objects.filter(subject=selected_subject, is_approved=False).select_related(
        'question_easy', 'question_medium', 'question_hard'
    ).order_by('-created_at')

    context = {
        'subjects': subjects,
        'selected_subject': selected_subject,
        'all_questions': all_questions,
        'exam_codes': exam_codes,
        'materials': materials,
        'pending_exam_codes': pending_exam_codes,
    }
    return render(request, 'qna/lecturer/lecturer_exam_codes_management.html', context)


@login_required
@require_POST
def lecturer_create_question(request):
    if not request.user.userprofile.is_lecturer:
        return JsonResponse({'success': False, 'error': 'Không có quyền truy cập'}, status=403)

    subject_id = request.POST.get('subject_id')
    question_text = request.POST.get('question_text', '').strip()
    difficulty = request.POST.get('difficulty')

    if not all([subject_id, question_text, difficulty]):
        return JsonResponse({'success': False, 'error': 'Thiếu thông tin bắt buộc'}, status=400)

    subject = get_object_or_404(request.user.userprofile.subjects_taught, id=subject_id)

    question = Question.objects.create(
        subject=subject,
        question_text=question_text,
        difficulty=difficulty,
        question_id_in_barem=f'Q_{subject.subject_code}_{uuid4().hex[:8]}',
        is_supplementary=False
    )

    messages.success(request, 'Đã tạo câu hỏi thành công.')
    return JsonResponse({'success': True, 'question_id': question.id})


@login_required
@require_POST
def lecturer_update_question(request, question_id):
    if not request.user.userprofile.is_lecturer:
        return JsonResponse({'success': False, 'error': 'Không có quyền truy cập'}, status=403)

    question = get_object_or_404(Question, pk=question_id)
    if question.subject not in request.user.userprofile.subjects_taught.all():
        return JsonResponse({'success': False, 'error': 'Không có quyền truy cập môn học này'}, status=403)

    question_text = request.POST.get('question_text', '').strip()
    difficulty = request.POST.get('difficulty')

    if question_text: question.question_text = question_text
    if difficulty: question.difficulty = difficulty
    question.save()

    messages.success(request, 'Đã cập nhật câu hỏi thành công.')
    return JsonResponse({'success': True})


@login_required
@require_POST
def lecturer_delete_question(request, question_id):
    if not request.user.userprofile.is_lecturer:
        return JsonResponse({'success': False, 'error': 'Không có quyền truy cập'}, status=403)

    question = get_object_or_404(Question, pk=question_id)
    if question.subject not in request.user.userprofile.subjects_taught.all():
        return JsonResponse({'success': False, 'error': 'Không có quyền truy cập môn học này'}, status=403)

    question.delete()
    messages.success(request, 'Đã xoá câu hỏi thành công.')
    return JsonResponse({'success': True})


@login_required
@require_POST
def lecturer_import_questions(request):
    if not request.user.userprofile.is_lecturer:
        return JsonResponse({'success': False, 'error': 'Không có quyền truy cập'}, status=403)

    subject_id = request.POST.get('subject_id')
    file_obj = request.FILES.get('file')
    if not subject_id or not file_obj:
        return JsonResponse({'success': False, 'error': 'Thiếu thông tin'}, status=400)

    subject = get_object_or_404(request.user.userprofile.subjects_taught, id=subject_id)

    try:
        import csv
        decoded_file = file_obj.read().decode('utf-8')
        io_string = io.StringIO(decoded_file)
        reader = csv.DictReader(io_string)

        imported_count = 0
        for row in reader:
            question_text = row.get('question', '').strip()
            difficulty = row.get('difficulty', 'MEDIUM').strip()
            if question_text and difficulty in ['EASY', 'MEDIUM', 'HARD']:
                Question.objects.create(
                    subject=subject,
                    question_text=question_text,
                    difficulty=difficulty,
                    question_id_in_barem=f'IMP_{subject.subject_code}_{uuid4().hex[:8]}',
                    is_supplementary=False
                )
                imported_count += 1

        messages.success(request, f'Đã import {imported_count} câu hỏi thành công.')
        return JsonResponse({'success': True, 'imported_count': imported_count})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)

# ===========================
# QUESTION BANK MANAGEMENT - gộp trực tiếp vào views.py
# ===========================

QUESTION_JOB_PREFIX = "qna_question_job_"


def _ensure_lecturer(request):
    if not request.user.userprofile.is_lecturer:
        raise PermissionDenied("Không có quyền truy cập.")


def _get_lecturer_subjects(request):
    return request.user.userprofile.subjects_taught.all().order_by("name")


def _get_selected_subject_for_lecturer(request, subject_id):
    subjects = _get_lecturer_subjects(request)
    if subject_id:
        return get_object_or_404(subjects, id=subject_id)
    return subjects.first()


def _serialize_material(material: LectureMaterial):
    return {
        "document_id": material.id,
        "file_name": material.title,
        "original_file_name": material.filename,
        "upload_time": material.uploaded_at.strftime("%d/%m/%Y"),
        "content_type": material.file_type,
        "extension": material.extension,
        "file_size": 0,
    }


def _serialize_question(question: Question):
    # 1. Nhận diện nguồn gốc dựa vào tiền tố
    barem_id = question.question_id_in_barem or ""
    if "AI_" in barem_id:
        source_text = "AI Tạo tự động"
    elif "IMP_" in barem_id:
        source_text = "Import từ file CSV/Excel"
    else:
        source_text = "Giảng viên thêm thủ công"

    # 2. Format thời gian tạo (Giờ:Phút - Ngày/Tháng/Năm)
    from django.utils import timezone
    time_str = ""
    if question.created_at:
        local_time = timezone.localtime(question.created_at)
        time_str = local_time.strftime("%H:%M - %d/%m/%Y")

    # 3. Trả dữ liệu về cho Frontend
    return {
        "question_id": question.id,
        "content": question.question_text,
        "difficulty": question.difficulty,
        "difficulty_label": question.get_difficulty_display(),
        "source": source_text,
        "source_type": "AI" if "AI_" in barem_id else "MANUAL",
        "created_at": time_str,  # <--- DỮ LIỆU THỜI GIAN ĐÃ ĐƯỢC TRUYỀN VÀO ĐÂY
        "updated_at": "",
    }


def _build_level_config(total_count, preset):
    total_count = int(total_count)
    if preset == "EASY":
        return {"easy": total_count, "medium": 0, "hard": 0}
    if preset == "MEDIUM":
        return {"easy": 0, "medium": total_count, "hard": 0}
    if preset == "HARD":
        return {"easy": 0, "medium": 0, "hard": total_count}

    easy = total_count // 3
    medium = total_count // 3
    hard = total_count - easy - medium
    return {"easy": easy, "medium": medium, "hard": hard}


def _save_question_job(job_id, payload):
    cache.set(f"{QUESTION_JOB_PREFIX}{job_id}", payload, timeout=60 * 30)


def _get_question_job(job_id):
    return cache.get(f"{QUESTION_JOB_PREFIX}{job_id}")


def _run_fake_question_generation(job_id, subject_id, document_ids, total_count, level_config):
    try:
        subject = Subject.objects.get(id=subject_id)
        materials = list(LectureMaterial.objects.filter(subject=subject, id__in=document_ids))

        plan = (
            ["EASY"] * int(level_config.get("easy", 0)) +
            ["MEDIUM"] * int(level_config.get("medium", 0)) +
            ["HARD"] * int(level_config.get("hard", 0))
        )

        questions = []
        for idx, difficulty in enumerate(plan, start=1):
            material = materials[(idx - 1) % len(materials)] if materials else None
            material_name = material.title if material else subject.name

            if difficulty == "EASY":
                content = f"Hãy trình bày khái niệm cơ bản liên quan đến nội dung trong tài liệu '{material_name}'."
            elif difficulty == "MEDIUM":
                content = f"Hãy phân tích nội dung trọng tâm trong tài liệu '{material_name}' và cho ví dụ minh họa."
            else:
                content = f"Hãy vận dụng kiến thức trong tài liệu '{material_name}' để giải thích một tình huống thực tế."

            questions.append({
                "temp_id": f"tmp_{uuid4().hex[:12]}",
                "content": content,
                "difficulty": difficulty,
                "source": material_name,
                "created_at": timezone.now().strftime("%d/%m/%Y"),
            })

        summary = {
            "all": len(questions),
            "easy": len([q for q in questions if q["difficulty"] == "EASY"]),
            "medium": len([q for q in questions if q["difficulty"] == "MEDIUM"]),
            "hard": len([q for q in questions if q["difficulty"] == "HARD"]),
        }

        _save_question_job(job_id, {
            "status": "COMPLETE",
            "progress": 100,
            "questions": questions,
            "summary": summary,
            "error_message": "",
        })
    except Exception as exc:
        logger.exception("Question generation failed")
        _save_question_job(job_id, {
            "status": "FAIL",
            "progress": 100,
            "questions": [],
            "summary": {},
            "error_message": str(exc),
        })


@login_required
def lecturer_questions_screen(request):
    _ensure_lecturer(request)

    subjects = _get_lecturer_subjects(request)
    subject_id = request.GET.get("subject_id")
    page_mode = request.GET.get("mode", "list")
    view_type = request.GET.get("view", "bank")

    selected_subject = _get_selected_subject_for_lecturer(request, subject_id)
    question_banks = []
    documents = []

    if selected_subject:
        # ==============================================================
        # THÊM ĐOẠN NÀY: DỌN DẸP NHÁP CŨ KHI VÀO MÀN HÌNH TẠO CÂU HỎI
        # ==============================================================
        if page_mode == "detail" and view_type == "generate":
            Question.objects.filter(
                subject=selected_subject,
                question_id_in_barem__startswith="DRAFT_"
            ).delete()
        # ==============================================================

        # Chỉ đếm những câu đã lưu (SAVED_) hoặc câu cũ (không phải DRAFT)
        saved_count = Question.objects.filter(subject=selected_subject).exclude(question_id_in_barem__startswith="DRAFT_").count()

        if saved_count > 0:
            question_banks = [{
                "bank_id": selected_subject.id,
                "name": f"Ngân hàng câu hỏi - {selected_subject.subject_code}",
                "detail_url": f"{request.path}?subject_id={selected_subject.id}&mode=detail&view=bank",
                "question_count": saved_count,
            }]

        if page_mode == "detail":
            documents = LectureMaterial.objects.filter(subject=selected_subject).order_by("-uploaded_at")

    context = {
        "subjects": subjects,
        "selected_subject": selected_subject,
        "page_mode": page_mode,
        "view_type": view_type,
        "question_banks": question_banks,
        "documents": [
            {"id": d.id, "file_name": d.title, "upload_time": d.uploaded_at.strftime("%d/%m/%Y")}
            for d in documents
        ] if page_mode == 'detail' else [],
        "upload_entry_url": f"{request.path}?subject_id={selected_subject.id}&mode=detail&view=generate" if selected_subject else "#",
    }
    return render(request, "qna/lecturer/lecturer_question_management.html", context)

@login_required
def question_bank_list_screen(request):
    return lecturer_questions_screen(request)


@login_required
def question_bank_detail_screen(request):
    if "mode" not in request.GET:
        q = request.GET.copy()
        q["mode"] = "detail"
        return redirect(f"{request.path}?{q.urlencode()}")
    return lecturer_questions_screen(request)


@login_required
@require_GET
def api_get_lecturer_subjects(request):
    _ensure_lecturer(request)
    subjects = _get_lecturer_subjects(request)
    return JsonResponse({
        "status": "SUCCESS",
        "subjects": [
            {"subject_id": str(item.id), "subject_name": item.name}
            for item in subjects
        ],
    })


@login_required
@require_GET
def api_get_question_banks(request):
    _ensure_lecturer(request)

    subject_id = request.GET.get("subject_id")
    if not subject_id:
        return JsonResponse({"status": "FAIL", "message": "Thiếu subject_id."}, status=400)

    subject = get_object_or_404(_get_lecturer_subjects(request), id=subject_id)
    question_count = Question.objects.filter(subject=subject).count()

    return JsonResponse({
        "status": "SUCCESS",
        "question_banks": [
            {
                "bank_id": str(subject.id),
                "bank_name": f"Ngân hàng câu hỏi - {subject.subject_code}",
                "bank_type": "ORAL",
                "created_at": "",
                "question_count": question_count,
            }
        ],
    })


@login_required
@require_POST
def api_create_question_bank(request):
    _ensure_lecturer(request)

    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        payload = request.POST

    subject_id = payload.get("subject_id")
    if not subject_id:
        return JsonResponse({"status": "FAIL", "message": "Thiếu subject_id."}, status=400)

    subject = get_object_or_404(_get_lecturer_subjects(request), id=subject_id)
    bank_name = (payload.get("name") or f"Ngân hàng câu hỏi - {subject.subject_code}").strip()

    return JsonResponse({
        "status": "SUCCESS",
        "bank_id": str(subject.id),
        "bank_name": bank_name,
        "message": "Tạo ngân hàng câu hỏi thành công.",
    })


@login_required
@require_POST
def api_save_question_bank_questions(request, bank_id):
    _ensure_lecturer(request)
    subject = get_object_or_404(_get_lecturer_subjects(request), id=bank_id)

    try:
        payload = json.loads(request.body.decode("utf-8"))
        question_ids = payload.get("question_ids", [])
    except Exception:
        return JsonResponse({"status": "FAIL", "message": "Dữ liệu không hợp lệ."}, status=400)

    # 1. Chuyển các câu được chọn thành SAVED_
    for q in Question.objects.filter(id__in=question_ids, subject=subject):
        q.question_id_in_barem = q.question_id_in_barem.replace("DRAFT_", "SAVED_")
        q.save()

    # 2. Xóa tất cả các câu DRAFT_ không được chọn còn sót lại của môn này
    Question.objects.filter(subject=subject, question_id_in_barem__startswith="DRAFT_").delete()

    return JsonResponse({
        "status": "SUCCESS",
        "subject_id": subject.id,
        "message": "Lưu ngân hàng câu hỏi thành công.",
    })


@login_required
@require_POST
def api_material_presign(request):
    """
    Presign giả cho môi trường local/dev.
    FE vẫn gọi bước này trước, nhưng upload thật sẽ hoàn tất ở api_material_upload_complete.
    """
    if not request.user.userprofile.is_lecturer:
        return JsonResponse({"status": "FAIL", "message": "Không có quyền truy cập."}, status=403)

    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        payload = request.POST

    file_name = (payload.get("file_name") or "").strip()
    file_size = int(payload.get("file_size") or 0)
    content_type = (payload.get("content_type") or "").strip()

    if not file_name:
        return JsonResponse({"status": "FAIL", "message": "Tên file không hợp lệ."}, status=400)

    if file_size <= 0:
        return JsonResponse({"status": "FAIL", "message": "File được chọn không có dữ liệu."}, status=400)

    if file_size > 50 * 1024 * 1024:
        return JsonResponse(
            {"status": "FAIL", "message": "Dung lượng file vượt quá 50MB. Vui lòng kiểm tra lại."},
            status=400,
        )

    ext = file_name.split(".")[-1].lower() if "." in file_name else ""
    allowed_ext = {"pdf", "docx", "txt"}
    if ext not in allowed_ext:
        return JsonResponse({"status": "FAIL", "message": "Định dạng file không hợp lệ."}, status=400)

    # MIME mềm để tránh lỗi browser/mobile gửi khác nhau
    allowed_content_types = {
        "pdf": {"application/pdf"},
        "docx": {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/octet-stream",
        },
        "txt": {"text/plain", "application/octet-stream"},
    }
    if content_type and content_type not in allowed_content_types.get(ext, set()):
        # không chặn cứng hoàn toàn nếu FE/browser gửi mime lạ
        pass

    file_key = f"lecture_materials/tmp/{uuid4().hex}_{file_name}"

    return JsonResponse({
        "status": "SUCCESS",
        "presigned_url": f"/media/{file_key}",
        "file_key": file_key,
        "message": "Tạo phiên upload thành công.",
    })

@login_required
@require_POST
def api_material_upload_complete(request):
    _ensure_lecturer(request)

    import hashlib

    subject_id = request.POST.get("subject_id")
    title = request.POST.get("title") or request.POST.get("file_name")
    file_obj = request.FILES.get("file")

    if not subject_id or not title or not file_obj:
        return JsonResponse({"status": "FAIL", "message": "Thiếu dữ liệu upload."}, status=400)

    subject = get_object_or_404(_get_lecturer_subjects(request), id=subject_id)

    original_name = file_obj.name.strip()
    ext = original_name.split(".")[-1].lower() if "." in original_name else ""
    if ext not in {"pdf", "docx", "txt"}:
        return JsonResponse({"status": "FAIL", "message": "Định dạng file không hợp lệ."}, status=400)

    upload_dir = os.path.join(settings.MEDIA_ROOT, "lecture_materials", subject.subject_code)
    os.makedirs(upload_dir, exist_ok=True)

    # Đọc toàn bộ file để tính hash chống trùng
    file_bytes = file_obj.read()
    if not file_bytes:
        return JsonResponse({"status": "FAIL", "message": "File rỗng hoặc không hợp lệ."}, status=400)

    file_hash = hashlib.md5(file_bytes).hexdigest()
    file_size = len(file_bytes)

    # Quét trùng theo: cùng subject + cùng tên file, hoặc cùng nội dung hash
    existing_materials = LectureMaterial.objects.filter(subject=subject)

    for material in existing_materials:
        existing_path = getattr(material, "file_path", None)
        existing_name = os.path.basename(existing_path) if existing_path else ""

        same_name = existing_name.lower() == original_name.lower()

        same_hash = False
        if existing_path and os.path.exists(existing_path):
            try:
                with open(existing_path, "rb") as f:
                    existing_hash = hashlib.md5(f.read()).hexdigest()
                same_hash = existing_hash == file_hash
            except Exception:
                same_hash = False

        if same_name or same_hash:
            return JsonResponse({
                "status": "FAIL",
                "message": f"Tài liệu '{original_name}' đã tồn tại hoặc bị trùng nội dung."
            }, status=400)

    # Ghi file
    safe_name = original_name
    base_name, extension = os.path.splitext(safe_name)
    file_path = os.path.join(upload_dir, safe_name)

    counter = 1
    while os.path.exists(file_path):
        safe_name = f"{base_name}_{counter}{extension}"
        file_path = os.path.join(upload_dir, safe_name)
        counter += 1

    with open(file_path, "wb+") as destination:
        destination.write(file_bytes)

    material = LectureMaterial.objects.create(
        subject=subject,
        title=os.path.splitext(title)[0],
        file_path=file_path,
        file_type=file_obj.content_type or "application/octet-stream",
    )

    return JsonResponse({
        "status": "SUCCESS",
        "document_id": str(material.id),
        "file_name": material.title,
        "message": "Upload tài liệu thành công.",
    })

@login_required
@require_POST
def lecturer_delete_material(request, material_id):
    material = get_object_or_404(LectureMaterial, id=material_id)

    if not request.user.userprofile.is_lecturer:
        return JsonResponse({
            "status": "FAIL",
            "message": "Không có quyền truy cập."
        }, status=403)

    if material.subject not in request.user.userprofile.subjects_taught.all():
        return JsonResponse({
            "status": "FAIL",
            "message": "Bạn không có quyền xóa tài liệu này."
        }, status=403)

    try:
        if material.file_path and os.path.exists(material.file_path):
            os.remove(material.file_path)
    except Exception as exc:
        logger.warning("Không thể xóa file vật lý: %s", exc)

    material.delete()

    return JsonResponse({
        "status": "SUCCESS",
        "message": "Đã xóa tài liệu thành công."
    })

@login_required
@require_GET
def api_get_materials(request):
    _ensure_lecturer(request)

    subject_id = request.GET.get("subject_id")
    search = (request.GET.get("search") or "").strip()

    if not subject_id:
        return JsonResponse({"status": "FAIL", "message": "Thiếu subject_id."}, status=400)

    subject = get_object_or_404(_get_lecturer_subjects(request), id=subject_id)
    qs = LectureMaterial.objects.filter(subject=subject).order_by("-uploaded_at")

    if search:
        qs = qs.filter(title__icontains=search)

    return JsonResponse({
        "status": "SUCCESS",
        "documents": [_serialize_material(item) for item in qs],
        "pagination": {
            "page": 1,
            "num_pages": 1,
            "has_next": False,
            "has_previous": False,
        },
    })


@login_required
@require_GET
def api_get_questions(request):
    _ensure_lecturer(request)

    bank_id = request.GET.get("bank_id")
    difficulty = request.GET.get("difficulty")

    # Lấy view_type, mặc định là bank nếu không truyền
    view_type = request.GET.get("view_type", "bank")

    if not bank_id:
        return JsonResponse({"status": "FAIL", "message": "Thiếu bank_id."}, status=400)

    subject = get_object_or_404(_get_lecturer_subjects(request), id=bank_id)
    qs = Question.objects.filter(subject=subject)

    # Nếu đang ở chế độ TẠO CÂU HỎI (generate) -> Chỉ lấy nháp (DRAFT_)
    if view_type == "generate":
        qs = qs.filter(question_id_in_barem__startswith="DRAFT_")
    # Nếu đang ở chế độ XEM NGÂN HÀNG (bank) -> Lấy TẤT CẢ trừ nháp
    else:
        qs = qs.exclude(question_id_in_barem__startswith="DRAFT_")

    if difficulty and difficulty != "ALL":
        qs = qs.filter(difficulty=difficulty)

    return JsonResponse({
        "status": "SUCCESS",
        "questions": [_serialize_question(item) for item in qs.order_by("-id")],
        "summary": {
            "all": qs.count(),
            "easy": qs.filter(difficulty=DifficultyLevel.EASY).count(),
            "medium": qs.filter(difficulty=DifficultyLevel.MEDIUM).count(),
            "hard": qs.filter(difficulty=DifficultyLevel.HARD).count(),
        },
    })


@login_required
@require_POST
def api_create_manual_question(request):
    _ensure_lecturer(request)

    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        payload = request.POST

    subject_id = payload.get("subject_id") or payload.get("bank_id")
    form = LecturerManualQuestionForm({
        "subject_id": subject_id,
        "question_text": payload.get("content") or payload.get("question_text"),
        "difficulty": payload.get("difficulty"),
    })

    if not form.is_valid():
        return JsonResponse({"status": "FAIL", "message": form.errors.as_json()}, status=400)

    subject = get_object_or_404(_get_lecturer_subjects(request), id=form.cleaned_data["subject_id"])

    question = Question.objects.create(
        subject=subject,
        question_text=form.cleaned_data["question_text"],
        difficulty=form.cleaned_data["difficulty"],
        question_id_in_barem=f"MAN_{subject.subject_code}_{uuid4().hex[:8]}",
        is_supplementary=False,
    )

    return JsonResponse({
        "status": "SUCCESS",
        "question": _serialize_question(question),
        "message": "Thêm câu hỏi thành công.",
    })


@login_required
@require_POST
def api_update_question_bank_question(request, question_id):
    _ensure_lecturer(request)

    question = get_object_or_404(Question, pk=question_id)
    if question.subject not in request.user.userprofile.subjects_taught.all():
        return JsonResponse({"status": "FAIL", "message": "Không có quyền truy cập môn học này."}, status=403)

    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        payload = request.POST

    question_text = (payload.get("content") or payload.get("question_text") or "").strip()
    difficulty = payload.get("difficulty")

    if question_text:
        question.question_text = question_text
    if difficulty in ["EASY", "MEDIUM", "HARD"]:
        question.difficulty = difficulty

    question.save()

    return JsonResponse({
        "status": "SUCCESS",
        "question": _serialize_question(question),
        "message": "Cập nhật câu hỏi thành công.",
    })


@login_required
@require_POST
def api_delete_question_bank_question(request, question_id):
    _ensure_lecturer(request)

    question = get_object_or_404(Question, pk=question_id)
    if question.subject not in request.user.userprofile.subjects_taught.all():
        return JsonResponse({"status": "FAIL", "message": "Không có quyền truy cập môn học này."}, status=403)

    question.delete()
    return JsonResponse({"status": "SUCCESS", "message": "Xóa câu hỏi thành công."})


@login_required
@require_POST
def api_bulk_update_question_level(request):
    _ensure_lecturer(request)

    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        payload = request.POST

    bank_id = payload.get("bank_id")
    difficulty = payload.get("difficulty")
    question_ids = payload.get("question_ids", [])

    if isinstance(question_ids, str):
        question_ids = json.loads(question_ids)

    subject = get_object_or_404(_get_lecturer_subjects(request), id=bank_id)

    updated = Question.objects.filter(
        subject=subject,
        id__in=question_ids,
    ).update(difficulty=difficulty)

    return JsonResponse({
        "status": "SUCCESS",
        "affected": updated,
        "message": "Thay đổi mức độ thành công.",
    })


@login_required
@require_POST
def api_bulk_delete_questions(request):
    _ensure_lecturer(request)

    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        payload = request.POST

    bank_id = payload.get("bank_id")
    question_ids = payload.get("question_ids", [])

    if isinstance(question_ids, str):
        question_ids = json.loads(question_ids)

    subject = get_object_or_404(_get_lecturer_subjects(request), id=bank_id)

    qs = Question.objects.filter(subject=subject, id__in=question_ids)
    deleted_count = qs.count()
    qs.delete()

    return JsonResponse({
        "status": "SUCCESS",
        "affected": deleted_count,
        "message": "Xóa câu hỏi thành công.",
    })


@login_required
@require_POST
def api_generate_questions(request):
    _ensure_lecturer(request)

    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        payload = request.POST

    subject_id = payload.get("subject_id")
    document_ids = payload.get("document_ids", [])
    total_count = int(payload.get("total_count", 0))
    level_config = payload.get("level_config", {})

    if isinstance(document_ids, str):
        document_ids = json.loads(document_ids)
    if isinstance(level_config, str):
        level_config = json.loads(level_config)

    if not subject_id:
        return JsonResponse({"status": "FAIL", "message": "Thiếu subject_id."}, status=400)

    subject = get_object_or_404(_get_lecturer_subjects(request), id=subject_id)

    if not document_ids:
        return JsonResponse({"status": "FAIL", "message": "Bạn phải chọn ít nhất 1 tài liệu."}, status=400)

    if total_count <= 0 or total_count > 100:
        return JsonResponse({"status": "FAIL", "message": "Số lượng câu hỏi không hợp lệ."}, status=400)

    materials_exist = LectureMaterial.objects.filter(subject=subject, id__in=document_ids).exists()
    if not materials_exist:
        return JsonResponse({"status": "FAIL", "message": "Danh sách tài liệu không hợp lệ."}, status=400)

    job_id = uuid4().hex
    _save_question_job(job_id, {
        "status": "PENDING",
        "progress": 0,
        "questions": [],
        "summary": {},
        "error_message": "",
    })

    channel_layer = get_channel_layer()
    async_to_sync(channel_layer.send)(
        "asr-tasks",
        {
            "type": "ai.generate_questions",
            # reply_channel giả để worker hiện tại không bị continue
            "reply_channel": f"question-job-{job_id}",
            "job_id": job_id,
            "subject_id": subject.id,
            "document_ids": document_ids,
            "total_count": total_count,
            "level_config": level_config,
        },
    )

    return JsonResponse({
        "status": "SUCCESS",
        "job_id": job_id,
        "message": "Đã tiếp nhận yêu cầu tạo câu hỏi.",
    }, status=202)


@login_required
@require_GET
def api_generate_questions_status(request):
    _ensure_lecturer(request)

    job_id = request.GET.get("job_id")
    if not job_id:
        return JsonResponse({"status": "FAIL", "message": "Thiếu job_id."}, status=400)

    job = _get_question_job(job_id)
    if not job:
        return JsonResponse({"status": "FAIL", "message": "Không tìm thấy job."}, status=404)

    return JsonResponse({
        "status": "SUCCESS",
        "process_status": job.get("status"),
        "progress": job.get("progress", 0),
        "questions": job.get("questions", []),
        "summary": job.get("summary", {}),
        "error_message": job.get("error_message", ""),
    })


@login_required
def lecturer_generate_codes_screen(request):
    if not request.user.userprofile.is_lecturer:
        return redirect('qna:dashboard')

    subjects = request.user.userprofile.subjects_taught.all()
    selected_subject_id = request.GET.get('subject_id')

    if not selected_subject_id:
        selected_subject = subjects.first()
    else:
        try:
            selected_subject = subjects.get(id=selected_subject_id)
        except (Subject.DoesNotExist, ValueError):
            selected_subject = subjects.first()

    materials = LectureMaterial.objects.filter(subject=selected_subject).order_by('-uploaded_at')
    pending_exam_codes = ExamCode.objects.filter(subject=selected_subject, is_approved=False).select_related(
        'question_easy', 'question_medium', 'question_hard'
    ).order_by('-created_at')
    approved_exam_codes = ExamCode.objects.filter(subject=selected_subject, is_approved=True).select_related(
        'question_easy', 'question_medium', 'question_hard'
    ).order_by('-created_at')

    context = {
        'subjects': subjects,
        'selected_subject': selected_subject,
        'materials': materials,
        'pending_exam_codes': pending_exam_codes,
        'approved_exam_codes': approved_exam_codes,
    }
    return render(request, 'qna/lecturer/lecturer_generate_exam_codes.html', context)


@login_required
@require_POST
def lecturer_upload_material_screen(request):
    if not request.user.userprofile.is_lecturer:
        return JsonResponse({'success': False, 'error': 'Không có quyền truy cập'}, status=403)

    subject_id = request.POST.get('subject_id')
    title = request.POST.get('title') or request.POST.get('material_name')  # Hỗ trợ cả 2 tên form
    file_obj = request.FILES.get('file') or request.FILES.get('material_file')

    if not subject_id or not title or not file_obj:
        return JsonResponse({'success': False, 'error': 'Thiếu thông tin bắt buộc'}, status=400)

    subject = get_object_or_404(request.user.userprofile.subjects_taught, id=subject_id)

    import os
    from django.conf import settings
    upload_dir = os.path.join(settings.MEDIA_ROOT, 'lecture_materials', subject.subject_code)
    os.makedirs(upload_dir, exist_ok=True)

    file_path = os.path.join(upload_dir, file_obj.name)
    with open(file_path, 'wb+') as destination:
        for chunk in file_obj.chunks():
            destination.write(chunk)

    material = LectureMaterial.objects.create(
        subject=subject,
        title=title,
        file_path=file_path,
        file_type=file_obj.content_type
    )

    return JsonResponse({'success': True, 'material_id': material.id})


@login_required
@require_POST
def lecturer_generate_codes_with_ai(request):
    if not request.user.userprofile.is_lecturer:
        return JsonResponse({'success': False, 'error': 'Không có quyền truy cập'}, status=403)

    subject_id = request.POST.get('subject_id')
    material_id = request.POST.get('material_id')
    num_codes = int(request.POST.get('num_codes', 4))

    subject = get_object_or_404(request.user.userprofile.subjects_taught, id=subject_id)
    if num_codes < 1 or num_codes > 16:
        return JsonResponse({'success': False, 'error': 'Số lượng mã đề phải từ 1 đến 16'}, status=400)

    material = get_object_or_404(LectureMaterial, pk=material_id, subject=subject)

    try:
        import openai
        client = openai.OpenAI()

        with open(material.file_path, 'r', encoding='utf-8') as f:
            content = f.read()

        system_prompt = """Bạn là chuyên gia giáo dục, chuyên tạo câu hỏi thi vấn đáp.
        ĐẦU RA: JSON format với 3 câu hỏi: {"easy": "...", "medium": "...", "hard": "..."}"""

        user_prompt = f"Tạo {num_codes} bộ câu hỏi (mỗi bộ 3 câu hỏi) cho môn học: {subject.name}\nNỘI DUNG:\n{content[:5000]}"

        resp = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt}
            ],
            temperature=0.7,
            response_format={"type": "json_object"}
        )

        questions_list = json.loads(resp.choices[0].message.content)
        created_codes = []

        # Nếu AI trả về dict đơn thay vì list, ta bọc lại thành list
        if isinstance(questions_list, dict) and "easy" in questions_list:
            questions_list = [questions_list]

        for i, questions in enumerate(questions_list):
            q_easy = Question.objects.create(
                subject=subject, question_text=questions['easy'],
                question_id_in_barem=f'AI_EASY_{uuid4().hex[:8]}', difficulty='EASY'
            )
            q_medium = Question.objects.create(
                subject=subject, question_text=questions['medium'],
                question_id_in_barem=f'AI_MEDIUM_{uuid4().hex[:8]}', difficulty='MEDIUM'
            )
            q_hard = Question.objects.create(
                subject=subject, question_text=questions['hard'],
                question_id_in_barem=f'AI_HARD_{uuid4().hex[:8]}', difficulty='HARD'
            )

            exam_code = ExamCode.objects.create(
                subject=subject, code_name=f'Mã đề AI-{uuid4().hex[:4]}',
                question_easy=q_easy, question_medium=q_medium, question_hard=q_hard,
                source_material=material.title, is_approved=False
            )
            created_codes.append(exam_code.id)

        return JsonResponse({'success': True, 'created_count': len(created_codes)})
    except Exception as e:
        logger.error(f"Lỗi khi sinh mã đề thi bằng AI: {e}")
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
def lecturer_create_session_screen(request):
    if not request.user.userprofile.is_lecturer:
        return redirect('qna:dashboard')

    subjects = request.user.userprofile.subjects_taught.all()
    selected_subject_id = request.GET.get('subject_id')

    if not selected_subject_id:
        selected_subject = subjects.first()
    else:
        try:
            selected_subject = subjects.get(id=selected_subject_id)
        except (Subject.DoesNotExist, ValueError):
            selected_subject = subjects.first()

    approved_exam_codes = ExamCode.objects.filter(subject=selected_subject, is_approved=True)
    all_rooms = ExamRoom.objects.all()
    exam_groups = ExamSessionGroup.objects.filter(subject=selected_subject).prefetch_related('rooms',
                                                                                             'exam_codes').order_by(
        '-exam_date')

    context = {
        'subjects': subjects,
        'selected_subject': selected_subject,
        'approved_exam_codes': approved_exam_codes,
        'all_rooms': all_rooms,
        'exam_groups': exam_groups,
    }
    return render(request, 'qna/lecturer/lecturer_create_exam_session.html', context)


@login_required
@require_POST
def lecturer_create_exam_group_screen(request):
    if not request.user.userprofile.is_lecturer:
        return JsonResponse({'success': False, 'error': 'Không có quyền truy cập'}, status=403)

    subject_id = request.POST.get('subject_id')
    group_name = request.POST.get('group_name')
    exam_date_str = request.POST.get('exam_date')
    duration_minutes = int(request.POST.get('duration_minutes', 60))
    exam_password = request.POST.get('exam_password', '').strip()
    exam_code_ids = request.POST.getlist('exam_code_ids')
    room_ids = request.POST.getlist('room_ids')

    subject = get_object_or_404(request.user.userprofile.subjects_taught, id=subject_id)
    if not group_name or not exam_date_str:
        return JsonResponse({'success': False, 'error': 'Thiếu thông tin bắt buộc'}, status=400)

    from datetime import datetime
    exam_date = datetime.strptime(exam_date_str, '%Y-%m-%dT%H:%M')

    exam_group = ExamSessionGroup.objects.create(
        subject=subject, group_name=group_name, exam_date=exam_date,
        duration_minutes=duration_minutes, exam_password=exam_password if exam_password else None,
        status='SCHEDULED', created_by=request.user
    )

    if exam_code_ids:
        exam_group.exam_codes.set(ExamCode.objects.filter(id__in=exam_code_ids, subject=subject))

    if room_ids:
        for room_id in room_ids:
            ExamSessionRoom.objects.create(exam_group=exam_group, room_id=room_id)

    messages.success(request, 'Đã tạo ca thi thành công.')
    return JsonResponse({'success': True, 'exam_group_id': exam_group.id})


@login_required
def lecturer_exam_sessions_list(request):
    if not request.user.userprofile.is_lecturer:
        return redirect('qna:dashboard')

    subjects = request.user.userprofile.subjects_taught.all()
    selected_subject_id = request.GET.get('subject_id')
    status_filter = request.GET.get('status')

    if not selected_subject_id:
        selected_subject = subjects.first()
    else:
        try:
            selected_subject = subjects.get(id=selected_subject_id)
        except (Subject.DoesNotExist, ValueError):
            selected_subject = subjects.first()

    exam_groups = ExamSessionGroup.objects.filter(subject=selected_subject).prefetch_related('rooms',
                                                                                             'exam_codes').order_by(
        '-exam_date')
    if status_filter:
        exam_groups = exam_groups.filter(status=status_filter)

    return render(request, 'qna/lecturer/lecturer_exam_sessions_list.html', {
        'subjects': subjects, 'selected_subject': selected_subject,
        'exam_groups': exam_groups, 'selected_status': status_filter,
        'status_choices': ExamSessionGroup.STATUS_CHOICES,
    })


@login_required
def lecturer_student_review_screen(request):
    if not request.user.userprofile.is_lecturer:
        return redirect('qna:dashboard')

    subjects = request.user.userprofile.subjects_taught.all()
    selected_subject_id = request.GET.get('subject_id')
    exam_group_id = request.GET.get('exam_group_id')
    student_filter = request.GET.get('student', '').strip()

    if not selected_subject_id:
        selected_subject = subjects.first()
    else:
        try:
            selected_subject = subjects.get(id=selected_subject_id)
        except (Subject.DoesNotExist, ValueError):
            selected_subject = subjects.first()

    exam_groups = ExamSessionGroup.objects.filter(subject=selected_subject).order_by('-exam_date')
    sessions = ExamSession.objects.filter(subject=selected_subject).select_related('user', 'user__userprofile')

    if exam_group_id: sessions = sessions.filter(exam_group_id=exam_group_id)
    if student_filter: sessions = sessions.filter(user__username__icontains=student_filter)

    sessions = sessions.order_by('-created_at')
    for session in sessions:
        main_avg, supp_sum, final_total = _compute_scores(session)
        session.main_avg = main_avg
        session.supp_sum = supp_sum
        session.calculated_final_score = final_total

    return render(request, 'qna/lecturer/lecturer_student_review.html', {
        'subjects': subjects, 'selected_subject': selected_subject,
        'exam_groups': exam_groups, 'selected_exam_group_id': int(exam_group_id) if exam_group_id else None,
        'sessions': sessions, 'student_filter': student_filter,
    })


@login_required
def lecturer_export_reports_screen(request):
    return redirect('qna:lecturer_student_review_screen')


@login_required
def lecturer_export_exam_results_screen(request):
    if not request.user.userprofile.is_lecturer:
        return redirect('qna:dashboard')

    subject_id = request.GET.get('subject_id')
    exam_group_id = request.GET.get('exam_group_id')
    subject = get_object_or_404(request.user.userprofile.subjects_taught, id=subject_id)

    if exam_group_id:
        exam_group = get_object_or_404(ExamSessionGroup, pk=exam_group_id, subject=subject)
        sessions = ExamSession.objects.filter(subject=subject, exam_group=exam_group).select_related('user',
                                                                                                     'user__userprofile')
    else:
        sessions = ExamSession.objects.filter(subject=subject).select_related('user', 'user__userprofile')

    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from datetime import datetime

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kết quả thi"

    headers = ['STT', 'Mã SV', 'Họ tên', 'Lớp', 'Ngày thi', 'Điểm chính', 'Điểm phụ', 'Điểm tổng']
    ws.append(headers)

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for idx, session in enumerate(sessions, 1):
        main_avg, supp_sum, final_total = _compute_scores(session)
        profile = session.user.userprofile
        ws.append([
            idx, session.user.username, profile.full_name, profile.class_name,
            session.created_at.strftime('%d/%m/%Y %H:%M'),
            f"{main_avg:.2f}", f"{supp_sum:.2f}", f"{final_total:.2f}"
        ])

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length: max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f'Ket_qua_thi_{subject.subject_code}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
def lecturer_subject_dashboard(request, subject_code):
    if not request.user.userprofile.is_lecturer:
        return redirect('qna:dashboard')

    subject = get_object_or_404(request.user.userprofile.subjects_taught, subject_code=subject_code)
    sessions_in_subject = ExamSession.objects.filter(subject=subject)

    context = {
        'subject': subject,
        'total_exams': sessions_in_subject.count(),
        'total_students': sessions_in_subject.values('user').distinct().count(),
        'completed_exams': sessions_in_subject.filter(is_completed=True).count(),
        'recent_exams': sessions_in_subject.select_related('user', 'user__userprofile').order_by('-created_at')[:10],
    }
    return render(request, 'qna/lecturer/lecturer_subject_dashboard.html', context)


@login_required
def update_exam_password(request, subject_code):
    if not request.user.userprofile.is_lecturer:
        return redirect('qna:dashboard')

    if request.method == 'POST':
        subject = get_object_or_404(request.user.userprofile.subjects_taught, subject_code=subject_code)
        password = request.POST.get('exam_password', '').strip()
        subject.exam_password = password if password else None
        subject.save()
        messages.success(request, 'Đã cập nhật mật khẩu bài thi thành công.')

    return redirect('qna:lecturer_subject_dashboard', subject_code=subject_code)


@login_required
def exam_password_view(request, subject_code):
    subject = get_object_or_404(Subject, subject_code=subject_code)
    if not subject.exam_password:
        return redirect('qna:pre_exam_verification', subject_code=subject_code)
    return render(request, 'qna/student/exam_password.html', {'subject': subject, 'subject_code': subject_code})


@login_required
def verify_exam_password(request, subject_code):
    if request.method == 'POST':
        subject = get_object_or_404(Subject, subject_code=subject_code)
        password = request.POST.get('password', '').strip()
        if not subject.exam_password or password == subject.exam_password:
            return redirect('qna:pre_exam_verification', subject_code=subject_code)
        else:
            messages.error(request, 'Mật khẩu không đúng. Vui lòng thử lại.')
            return redirect('qna:exam_password', subject_code=subject_code)
    return redirect('qna:exam_password', subject_code=subject_code)


SUPP_MAX_PER_QUESTION = 1.0
SUPP_MAX_COUNT = 2
FINAL_CAP = 7.0


def _json_body(request: HttpRequest) -> Dict[str, Any]:
    try:
        if request.body: return json.loads(request.body.decode("utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError):
        pass
    return {}


def _ensure_owner(session: ExamSession, user: User) -> None:
    if session.user_id != getattr(user, "id", None):
        raise PermissionDenied("Bạn không có quyền truy cập phiên thi này.")


def _compute_scores(session: ExamSession) -> tuple[float, float, float]:
    total_main_questions = session.questions.count() or 3
    main_avg = (ExamResult.objects.filter(session=session).aggregate(total=Sum('score'))[
                    'total'] or 0.0) / total_main_questions

    supp_results_qs = SupplementaryResult.objects.filter(session=session)
    best_scores_by_text = {}
    for result in supp_results_qs:
        key = (result.question_text or "").strip()
        if not key: continue
        score = float(result.score or 0.0)
        if score > SUPP_MAX_PER_QUESTION: score /= 10.0
        current_score = max(0.0, min(score, SUPP_MAX_PER_QUESTION))
        if key not in best_scores_by_text or current_score > best_scores_by_text[key]:
            best_scores_by_text[key] = current_score

    unique_supp_scores = sorted(list(best_scores_by_text.values()), reverse=True)
    supp_sum = sum(unique_supp_scores[:SUPP_MAX_COUNT])
    final_total = min(FINAL_CAP, main_avg + supp_sum) if supp_sum > 0 else min(10.0, main_avg)

    return main_avg, supp_sum, final_total


def _dedupe_supp_for_display(qs: SupplementaryResult) -> List[SupplementaryResult]:
    best_by_text = {}
    for result in qs:
        key = (result.question_text or "").strip()
        if not key: continue
        cleaned_result = result
        score = float(cleaned_result.score or 0.0)
        if score > SUPP_MAX_PER_QUESTION: cleaned_result.score = score / 10.0
        if best_by_text.get(key) is None or float(cleaned_result.score or 0) > float(best_by_text.get(key).score or 0):
            best_by_text[key] = cleaned_result
    return sorted(list(best_by_text.values()), key=lambda x: float(x.score or 0), reverse=True)[:SUPP_MAX_COUNT]


class RegistrationForm(forms.Form):
    full_name = forms.CharField(label=mark_safe('Họ và tên <span class="text-red-500">*</span>'), max_length=150)
    username = forms.CharField(label=mark_safe('Tên đăng nhập <span class="text-red-500">*</span>'), max_length=150)
    class_name = forms.CharField(label=mark_safe('Lớp <span class="text-red-500">*</span>'), max_length=100)
    email = forms.EmailField(label='Email', required=False)
    faculty = forms.CharField(label='Khoa', required=False, max_length=150)
    password = forms.CharField(label=mark_safe('Mật khẩu <span class="text-red-500">*</span>'),
                               widget=forms.PasswordInput())
    password2 = forms.CharField(label=mark_safe('Nhập lại mật khẩu <span class="text-red-500">*</span>'),
                                widget=forms.PasswordInput())

    def clean_username(self):
        username = self.cleaned_data.get("username", "").strip()
        if not username: raise ValidationError("Tên đăng nhập là bắt buộc.")
        if User.objects.filter(username=username).exists(): raise ValidationError("Tên đăng nhập này đã tồn tại.")
        return username

    def clean(self):
        cleaned_data = super().clean()
        password, password2 = cleaned_data.get("password"), cleaned_data.get("password2")
        if password and password2 and password != password2:
            self.add_error("password2", "Mật khẩu nhập lại không khớp.")
        if password:
            try:
                validate_password(password, user=User(username=cleaned_data.get("username")))
            except ValidationError as e:
                self.add_error("password", e)
        return cleaned_data


@login_required
def dashboard_view(request: HttpRequest) -> HttpResponse:
    subjects = Subject.objects.all().order_by("name")
    recent_sessions = ExamSession.objects.filter(user=request.user).select_related("subject").order_by("-created_at")[
        :5]
    profile, _ = UserProfile.objects.get_or_create(user=request.user)
    return render(request, "qna/student/dashboard.html", {
        "subjects": subjects, "recent_sessions": recent_sessions,
        "full_name": profile.full_name or request.user.first_name,
    })


@login_required
def history_view(request: HttpRequest) -> HttpResponse:
    sessions = ExamSession.objects.filter(user=request.user).select_related("subject").order_by("-created_at")
    for s in sessions:
        main_avg, supp_sum, final_total = _compute_scores(s)
        s.main_avg, s.supp_sum, s.calculated_final_score = main_avg, supp_sum, final_total
    return render(request, "qna/student/history.html", {"sessions": sessions})


@login_required
def history_detail_view(request: HttpRequest, session_id: int) -> HttpResponse:
    session = get_object_or_404(ExamSession.objects.select_related("subject", "user"), pk=session_id)
    _ensure_owner(session, request.user)
    main_results = ExamResult.objects.filter(session=session).select_related("question").order_by("question_id")
    supp_results_display = _dedupe_supp_for_display(SupplementaryResult.objects.filter(session=session))
    main_avg, supp_sum, final_total = _compute_scores(session)
    return render(request, "qna/student/history_detail.html", {
        "session": session, "results": main_results, "supp_results": supp_results_display,
        "main_avg": main_avg, "supp_sum": supp_sum, "final_total": final_total,
    })


@login_required
def pre_exam_verification_view(request: HttpRequest, subject_code: str) -> HttpResponse:
    return render(request, 'qna/student/pre_exam_verification.html',
                  {'subject': get_object_or_404(Subject, subject_code=subject_code), 'subject_code': subject_code})


@login_required
def exam_view(request: HttpRequest, subject_code: str) -> HttpResponse:
    subject = get_object_or_404(Subject, subject_code=subject_code)
    main_questions = list(Question.objects.filter(subject=subject, is_supplementary=False).order_by("?")[:3])
    if not main_questions:
        messages.error(request, f"Môn {subject.name} chưa có câu hỏi. Vui lòng liên hệ quản trị viên.")
        return redirect("qna:dashboard")

    session = ExamSession.objects.create(user=request.user, subject=subject)
    session.questions.set(main_questions)
    barem = [{"id": q.id, "question": q.question_text} for q in
             Question.objects.filter(subject=subject, is_supplementary=False).exclude(
                 id__in=[q.id for q in main_questions])]
    return render(request, "qna/student/exam.html", {
        "subject": subject, "selected_questions": main_questions, "session": session,
        "barem_json": json.dumps(barem, ensure_ascii=False),
    })


def _get_avatar_data_url(profile: UserProfile) -> str:
    if profile.profile_image_blob:
        try:
            return f"data:{profile.profile_image_mime or 'image/jpeg'};base64,{b64encode(profile.profile_image_blob).decode('ascii')}"
        except Exception:
            pass
    return static("images/default_avatar.png")


@login_required
def profile_view(request: HttpRequest) -> HttpResponse:
    try:
        profile = request.user.userprofile
    except UserProfile.DoesNotExist:
        profile = None
    return render(request, 'qna/student/profile.html', {'user': request.user, 'profile': profile})


@login_required
@require_POST
def update_profile_image(request: HttpRequest) -> JsonResponse:
    file_obj = request.FILES.get('profile_image')
    if not file_obj: return JsonResponse({"success": False, "error": "Không tìm thấy file ảnh."}, status=400)
    if file_obj.size > 5 * 1024 * 1024: return JsonResponse(
        {"success": False, "error": "Kích thước ảnh không được vượt quá 5MB."}, status=400)

    content = file_obj.read()
    mime = file_obj.content_type
    profile, _ = UserProfile.objects.get_or_create(user=request.user)
    profile.profile_image_blob = content
    profile.profile_image_mime = mime
    profile.save()
    return JsonResponse({"success": True, "image_data_url": f"data:{mime};base64,{b64encode(content).decode('ascii')}"})


@login_required
@require_POST
def save_exam_result(request: HttpRequest) -> JsonResponse:
    data = _json_body(request)
    if not all([data.get("session_id"), data.get("question_id"), data.get("score") is not None]):
        return HttpResponseBadRequest("Thiếu các tham số bắt buộc (session_id, question_id, score).")

    session = get_object_or_404(ExamSession, pk=data.get("session_id"))
    _ensure_owner(session, request.user)
    question = get_object_or_404(Question, pk=data.get("question_id"), is_supplementary=False)

    result, created = ExamResult.objects.update_or_create(
        session=session, question=question, defaults={
            "transcript": data.get("transcript", ""), "score": float(data.get("score")),
            "feedback": data.get("feedback"), "analysis": data.get("analysis"), "answered_at": timezone.now(),
        }
    )
    return JsonResponse({"status": "ok", "created": created, "result_id": result.id})


@login_required
@require_POST
def get_supplementary_for_session(request: HttpRequest, session_id: int) -> JsonResponse:
    session = get_object_or_404(ExamSession.objects.select_related("subject"), pk=session_id)
    _ensure_owner(session, request.user)
    supp_pool = list(Question.objects.filter(subject=session.subject, is_supplementary=True))
    random.shuffle(supp_pool)
    return JsonResponse(
        {"status": "ok", "items": [{"id": q.id, "question": q.question_text} for q in supp_pool[:SUPP_MAX_COUNT]]})


@login_required
@require_POST
def save_supplementary_result(request: HttpRequest) -> JsonResponse:
    data = _json_body(request)
    if not all([data.get("session_id"), (data.get("question_text") or "").strip(), data.get("score") is not None]):
        return HttpResponseBadRequest("Thiếu tham số.")

    session = get_object_or_404(ExamSession, pk=data.get("session_id"))
    _ensure_owner(session, request.user)
    if SupplementaryResult.objects.filter(session=session).count() >= SUPP_MAX_COUNT:
        return JsonResponse({"status": "error", "message": f"Đã đạt số lượng câu hỏi phụ tối đa ({SUPP_MAX_COUNT})."},
                            status=400)

    try:
        max_score_val = float(
            data.get("max_score", 10.0)); max_score_val = 10.0 if max_score_val <= 0 else max_score_val
    except (TypeError, ValueError):
        return HttpResponseBadRequest("Giá trị điểm không hợp lệ.")

    final_score = max(0.0,
                      min((float(data.get("score")) / max_score_val) * SUPP_MAX_PER_QUESTION, SUPP_MAX_PER_QUESTION))
    sr = SupplementaryResult.objects.create(session=session, question_text=(data.get("question_text") or "").strip(),
                                            transcript=data.get("transcript", ""), score=final_score,
                                            max_score=SUPP_MAX_PER_QUESTION, feedback=data.get("feedback"),
                                            analysis=data.get("analysis"))
    main_avg, supp_sum, final_total = _compute_scores(session)
    return JsonResponse({"status": "ok", "supplementary_result_id": sr.id, "main_avg": main_avg, "supp_sum": supp_sum,
                         "final_total": final_total})


@login_required
@require_POST
def finalize_session_view(request: HttpRequest, session_id: int) -> JsonResponse:
    session = get_object_or_404(ExamSession, pk=session_id)
    _ensure_owner(session, request.user)
    _, _, total_score = _compute_scores(session)
    session.is_completed = True
    session.completed_at = timezone.now()
    session.final_score = total_score
    session.save(update_fields=["is_completed", "completed_at", "final_score"])
    return JsonResponse({"status": "success", "final_score": session.final_score})


@login_required
@require_POST
def verify_student_face(request: HttpRequest) -> JsonResponse:
    face_image_data = request.POST.get('face_image')
    subject_code = request.POST.get('subject_code')
    if not face_image_data: return JsonResponse({'status': 'error', 'message': 'Thiếu ảnh khuôn mặt.'}, status=400)
    if not subject_code: return JsonResponse({'status': 'error', 'message': 'Thiếu mã môn học.'}, status=400)
    try:
        subject = get_object_or_404(Subject, subject_code=subject_code)
        format, imgstr = face_image_data.split(';base64,')
        session = ExamSession.objects.create(
            user=request.user, subject=subject,
            face_image_blob=base64.b64decode(imgstr), face_image_mime=f"image/{format.split('/')[-1]}",
            verification_status='ALLOW'
        )
        return JsonResponse(
            {'status': 'success', 'session_id': session.id, 'message': 'Đã lưu ảnh xác thực thành công.'})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': f'Lỗi khi xử lý ảnh: {str(e)}'}, status=500)


@login_required
def get_verification_images(request: HttpRequest, session_id: int) -> HttpResponse:
    session = get_object_or_404(ExamSession, pk=session_id)
    if not (request.user.userprofile.is_lecturer or session.user == request.user): raise PermissionDenied(
        "Bạn không có quyền xem ảnh xác thực này.")

    face_data_url = f"data:{session.face_image_mime};base64,{b64encode(session.face_image_blob).decode('ascii')}" if session.face_image_blob and session.face_image_mime else None
    id_card_data_url = f"data:{session.id_card_image_mime};base64,{b64encode(session.id_card_image_blob).decode('ascii')}" if session.id_card_image_blob and session.id_card_image_mime else None

    return JsonResponse({
        'status': 'success', 'face_image': face_data_url, 'id_card_image': id_card_data_url,
        'verification_score': session.verification_score, 'verification_status': session.verification_status,
        'needs_manual_review': session.needs_manual_review
    })


@login_required
def lecturer_question_management(request, subject_code):
    """
    Màn quản lý câu hỏi theo subject_code, nhưng dùng cùng logic/context
    với giao diện question bank mới để không vỡ template.
    """
    if not request.user.userprofile.is_lecturer:
        messages.error(request, "Bạn không có quyền truy cập.")
        return redirect("qna:dashboard")

    subject = get_object_or_404(
        request.user.userprofile.subjects_taught.all(),
        subject_code=subject_code
    )

    subjects = request.user.userprofile.subjects_taught.all().order_by("name")
    page_mode = request.GET.get("mode", "detail")

    question_banks = [{
        "bank_id": subject.id,
        "name": f"Ngân hàng câu hỏi - {subject.subject_code}",
        "detail_url": f"{request.path}?mode=detail",
        "question_count": Question.objects.filter(subject=subject).count(),
    }]

    documents = LectureMaterial.objects.filter(subject=subject).order_by("-uploaded_at")

    questions = [
        {
            "id": q.id,
            "content": q.question_text,
            "level": q.difficulty,
            "source_name": "",
            "created_at": "",
            "selected": False,
        }
        for q in Question.objects.filter(subject=subject).order_by("question_id_in_barem", "-id")
    ]

    context = {
        "subject": subject,
        "subjects": subjects,
        "selected_subject": subject,
        "page_mode": page_mode,
        "question_banks": question_banks,
        "documents": [
            {
                "id": d.id,
                "file_name": d.title,
                "upload_time": d.uploaded_at.strftime("%d/%m/%Y"),
            }
            for d in documents
        ],
        "questions": questions,
        "upload_entry_url": request.path,
    }
    return render(request, "qna/lecturer/lecturer_question_management.html", context)

@login_required
@require_POST
def lecturer_update_exam_code_question(request, exam_code_id):
    if not request.user.userprofile.is_lecturer: return JsonResponse(
        {'success': False, 'error': 'Không có quyền truy cập'}, status=403)
    exam_code = get_object_or_404(ExamCode, pk=exam_code_id)
    if exam_code.subject not in request.user.userprofile.subjects_taught.all(): return JsonResponse(
        {'success': False, 'error': 'Không có quyền truy cập'}, status=403)

    difficulty = request.POST.get('difficulty')
    if difficulty not in ['EASY', 'MEDIUM', 'HARD']: return JsonResponse(
        {'success': False, 'error': 'Mức độ khó không hợp lệ'}, status=400)

    question = get_object_or_404(Question, pk=request.POST.get('question_id'),
                                 subject=exam_code.subject) if request.POST.get('question_id') else None
    if difficulty == 'EASY':
        exam_code.question_easy = question
    elif difficulty == 'MEDIUM':
        exam_code.question_medium = question
    elif difficulty == 'HARD':
        exam_code.question_hard = question
    exam_code.save()

    messages.success(request, 'Đã cập nhật câu hỏi thành công.')
    return JsonResponse({'success': True})


@login_required
@require_POST
def lecturer_delete_exam_code(request, exam_code_id):
    if not request.user.userprofile.is_lecturer: return JsonResponse(
        {'success': False, 'error': 'Không có quyền truy cập'}, status=403)
    exam_code = get_object_or_404(ExamCode, pk=exam_code_id)
    if exam_code.subject not in request.user.userprofile.subjects_taught.all(): return JsonResponse(
        {'success': False, 'error': 'Không có quyền truy cập'}, status=403)
    exam_code.delete()
    messages.success(request, 'Đã xoá mã đề thi thành công.')
    return JsonResponse({'success': True})


@login_required
def lecturer_generate_exam_codes(request, subject_code):
    if not request.user.userprofile.is_lecturer: return redirect('qna:dashboard')
    subject = get_object_or_404(request.user.userprofile.subjects_taught, subject_code=subject_code)
    return render(request, 'qna/lecturer/lecturer_generate_exam_codes.html', {
        'subject': subject, 'materials': LectureMaterial.objects.filter(subject=subject).order_by('-uploaded_at'),
        'pending_exam_codes': ExamCode.objects.filter(subject=subject, is_approved=False).select_related(
            'question_easy', 'question_medium', 'question_hard').order_by('-created_at'),
    })


@login_required
@require_POST
def lecturer_upload_material(request, subject_code):
    if not request.user.userprofile.is_lecturer: return JsonResponse({'success': False, 'error': 'Không có quyền'},
                                                                     status=403)
    subject = get_object_or_404(request.user.userprofile.subjects_taught, subject_code=subject_code)
    title, file_obj = request.POST.get('title'), request.FILES.get('file')
    if not title or not file_obj: return JsonResponse({'success': False, 'error': 'Thiếu thông tin'}, status=400)

    import os;
    from django.conf import settings
    upload_dir = os.path.join(settings.MEDIA_ROOT, 'lecture_materials', subject_code)
    os.makedirs(upload_dir, exist_ok=True)
    file_path = os.path.join(upload_dir, file_obj.name)
    with open(file_path, 'wb+') as destination:
        for chunk in file_obj.chunks(): destination.write(chunk)

    material = LectureMaterial.objects.create(subject=subject, title=title, file_path=file_path,
                                              file_type=file_obj.content_type)
    messages.success(request, 'Đã upload tài liệu thành công.')
    return JsonResponse({'success': True, 'material_id': material.id})


@login_required
@require_POST
def lecturer_approve_exam_code(request, exam_code_id):
    if not request.user.userprofile.is_lecturer: return JsonResponse({'success': False, 'error': 'Không có quyền'},
                                                                     status=403)
    exam_code = get_object_or_404(ExamCode, pk=exam_code_id)
    if exam_code.subject not in request.user.userprofile.subjects_taught.all(): return JsonResponse({'success': False},
                                                                                                    status=403)
    exam_code.is_approved = True;
    exam_code.save()
    messages.success(request, 'Đã duyệt mã đề thi thành công.')
    return JsonResponse({'success': True})


@login_required
@require_POST
def lecturer_edit_exam_code_question(request, exam_code_id):
    if not request.user.userprofile.is_lecturer: return JsonResponse({'success': False}, status=403)
    exam_code = get_object_or_404(ExamCode, pk=exam_code_id)
    if exam_code.subject not in request.user.userprofile.subjects_taught.all(): return JsonResponse({'success': False},
                                                                                                    status=403)
    difficulty, new_text = request.POST.get('difficulty'), request.POST.get('new_text', '').strip()
    if difficulty not in ['EASY', 'MEDIUM', 'HARD'] or not new_text: return JsonResponse(
        {'success': False, 'error': 'Lỗi dữ liệu'}, status=400)

    question = exam_code.question_easy if difficulty == 'EASY' else (
        exam_code.question_medium if difficulty == 'MEDIUM' else exam_code.question_hard)
    if question:
        question.question_text = new_text;
        question.save()
    else:
        question = Question.objects.create(subject=exam_code.subject, question_text=new_text,
                                           question_id_in_barem=f'AI_EDITED_{difficulty}_{uuid4().hex[:8]}',
                                           is_supplementary=False)
        if difficulty == 'EASY':
            exam_code.question_easy = question
        elif difficulty == 'MEDIUM':
            exam_code.question_medium = question
        elif difficulty == 'HARD':
            exam_code.question_hard = question
        exam_code.save()
    messages.success(request, 'Đã chỉnh sửa câu hỏi thành công.')
    return JsonResponse({'success': True})


@login_required
def lecturer_create_exam_session(request, subject_code):
    if not request.user.userprofile.is_lecturer: return redirect('qna:dashboard')
    subject = get_object_or_404(request.user.userprofile.subjects_taught, subject_code=subject_code)
    return render(request, 'qna/lecturer/lecturer_create_exam_session.html', {
        'subject': subject,
        'approved_exam_codes': ExamCode.objects.filter(subject=subject, is_approved=True).select_related(
            'question_easy', 'question_medium', 'question_hard'),
        'all_rooms': ExamRoom.objects.all(),
        'exam_groups': ExamSessionGroup.objects.filter(subject=subject).prefetch_related('rooms',
                                                                                         'exam_codes').order_by(
            '-exam_date')
    })


@login_required
@require_POST
def lecturer_create_exam_group(request, subject_code):
    if not request.user.userprofile.is_lecturer: return JsonResponse({'success': False}, status=403)
    subject = get_object_or_404(request.user.userprofile.subjects_taught, subject_code=subject_code)
    group_name, exam_date_str = request.POST.get('group_name'), request.POST.get('exam_date')
    if not group_name or not exam_date_str: return JsonResponse({'success': False, 'error': 'Thiếu thông tin'},
                                                                status=400)

    from datetime import datetime
    exam_group = ExamSessionGroup.objects.create(
        subject=subject, group_name=group_name, exam_date=datetime.strptime(exam_date_str, '%Y-%m-%dT%H:%M'),
        duration_minutes=int(request.POST.get('duration_minutes', 60)),
        exam_password=request.POST.get('exam_password', '').strip() or None,
        status='SCHEDULED', created_by=request.user
    )
    if request.POST.getlist('exam_code_ids'): exam_group.exam_codes.set(
        ExamCode.objects.filter(id__in=request.POST.getlist('exam_code_ids'), subject=subject))
    for room_id in request.POST.getlist('room_ids'): ExamSessionRoom.objects.create(exam_group=exam_group,
                                                                                    room_id=room_id)
    messages.success(request, 'Đã tạo ca thi thành công.')
    return JsonResponse({'success': True, 'exam_group_id': exam_group.id})


@login_required
@require_POST
def lecturer_update_exam_group(request, exam_group_id):
    if not request.user.userprofile.is_lecturer: return JsonResponse({'success': False}, status=403)
    exam_group = get_object_or_404(ExamSessionGroup, pk=exam_group_id)
    if exam_group.subject not in request.user.userprofile.subjects_taught.all() or exam_group.status == 'COMPLETED': return JsonResponse(
        {'success': False}, status=403)

    if 'group_name' in request.POST: exam_group.group_name = request.POST['group_name']
    if 'exam_date' in request.POST:
        from datetime import datetime;
        exam_group.exam_date = datetime.strptime(request.POST['exam_date'], '%Y-%m-%dT%H:%M')
    if 'duration_minutes' in request.POST: exam_group.duration_minutes = int(request.POST['duration_minutes'])
    if 'exam_password' in request.POST: exam_group.exam_password = request.POST['exam_password'].strip() or None
    exam_group.save()
    messages.success(request, 'Đã cập nhật ca thi thành công.')
    return JsonResponse({'success': True})


@login_required
@require_POST
def lecturer_import_students_to_room(request, session_room_id):
    if not request.user.userprofile.is_lecturer: return JsonResponse({'success': False}, status=403)
    session_room = get_object_or_404(ExamSessionRoom, pk=session_room_id)
    if session_room.exam_group.subject not in request.user.userprofile.subjects_taught.all(): return JsonResponse(
        {'success': False}, status=403)
    if not request.FILES.get('csv_file'): return JsonResponse({'success': False, 'error': 'Thiếu file'}, status=400)

    import csv, io
    try:
        reader, imported_count = csv.DictReader(io.StringIO(request.FILES.get('csv_file').read().decode('utf-8'))), 0
        for row in reader:
            if username := row.get('username', '').strip():
                try:
                    session_room.students.add(User.objects.get(username=username)); imported_count += 1
                except User.DoesNotExist:
                    continue
        messages.success(request, f'Đã import {imported_count} sinh viên thành công.')
        return JsonResponse({'success': True, 'imported_count': imported_count})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_POST
def lecturer_random_assign_students(request, exam_group_id):
    if not request.user.userprofile.is_lecturer: return JsonResponse({'success': False}, status=403)
    exam_group = get_object_or_404(ExamSessionGroup, pk=exam_group_id)
    if exam_group.subject not in request.user.userprofile.subjects_taught.all(): return JsonResponse({'success': False},
                                                                                                     status=403)

    all_students, session_rooms = [], list(exam_group.session_rooms.all())
    for sr in session_rooms: all_students.extend(list(sr.students.all())); sr.students.clear()

    random.shuffle(all_students)
    for i, student in enumerate(all_students): session_rooms[i % len(session_rooms)].students.add(student)
    messages.success(request, 'Đã random sinh viên vào các phòng thi thành công.')
    return JsonResponse({'success': True})


@login_required
def lecturer_student_review(request, subject_code):
    if not request.user.userprofile.is_lecturer: return redirect('qna:dashboard')
    subject = get_object_or_404(request.user.userprofile.subjects_taught, subject_code=subject_code)
    return render(request, 'qna/lecturer/lecturer_student_review.html', {
        'subject': subject, 'sessions': ExamSession.objects.filter(subject=subject).select_related('user',
                                                                                                   'user__userprofile').prefetch_related(
            'results__question').order_by('-created_at')
    })


@login_required
def lecturer_session_detail(request, session_id):
    if not request.user.userprofile.is_lecturer: return redirect('qna:dashboard')
    session = get_object_or_404(ExamSession.objects.select_related('subject', 'user', 'user__userprofile'),
                                pk=session_id)
    if session.subject not in request.user.userprofile.subjects_taught.all(): raise PermissionDenied(
        "Bạn không có quyền")

    main_avg, supp_sum, final_total = _compute_scores(session)
    return render(request, 'qna/lecturer/lecturer_session_detail.html', {
        'session': session,
        'results': ExamResult.objects.filter(session=session).select_related("question").order_by("question_id"),
        'supp_results': _dedupe_supp_for_display(SupplementaryResult.objects.filter(session=session)),
        'main_avg': main_avg, 'supp_sum': supp_sum, 'final_total': final_total,
    })


@login_required
def lecturer_export_exam_results(request, subject_code):
    if not request.user.userprofile.is_lecturer:
        messages.error(request, "Bạn không có quyền truy cập.")
        return redirect("qna:dashboard")

    subject = get_object_or_404(
        request.user.userprofile.subjects_taught.all(),
        subject_code=subject_code
    )

    sessions = (
        ExamSession.objects
        .filter(subject=subject, is_completed=True)
        .select_related("user", "exam_group")
        .order_by("-created_at")
    )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="ket_qua_{subject.subject_code}.xlsx"'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kết quả thi"

    headers = [
        "STT",
        "Họ tên",
        "Username",
        "Mã sinh viên",
        "Lớp",
        "Môn học",
        "Mã môn",
        "Ca thi",
        "Ngày thi",
        "Điểm cuối",
        "Trạng thái hoàn thành",
        "Xác thực khuôn mặt",
    ]
    ws.append(headers)

    for idx, session in enumerate(sessions, start=1):
        profile = getattr(session.user, "userprofile", None)
        ws.append([
            idx,
            profile.full_name if profile and profile.full_name else session.user.get_full_name() or session.user.username,
            session.user.username,
            profile.student_id if profile else "",
            profile.class_name if profile else "",
            subject.name,
            subject.subject_code,
            session.exam_group.group_name if session.exam_group else "",
            session.created_at.strftime("%d/%m/%Y %H:%M") if session.created_at else "",
            session.final_score if session.final_score is not None else "",
            "Đã hoàn thành" if session.is_completed else "Chưa hoàn thành",
            session.get_verification_status_display() if hasattr(session, "get_verification_status_display") else session.verification_status,
        ])

    # Auto width đơn giản
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            try:
                cell_length = len(str(cell.value)) if cell.value is not None else 0
                if cell_length > max_length:
                    max_length = cell_length
            except Exception:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 40)

    wb.save(response)
    return response


@login_required
def lecturer_manage_rooms(request):
    if not request.user.userprofile.is_lecturer: return redirect('qna:dashboard')
    return render(request, 'qna/lecturer/lecturer_manage_rooms.html',
                  {'rooms': ExamRoom.objects.all().order_by('room_code')})


@login_required
@require_POST
def lecturer_create_room(request):
    if not request.user.userprofile.is_lecturer: return JsonResponse({'success': False}, status=403)
    room_name, room_code = request.POST.get('room_name'), request.POST.get('room_code')
    if not room_name or not room_code: return JsonResponse({'success': False, 'error': 'Thiếu thông tin'}, status=400)
    room = ExamRoom.objects.create(room_name=room_name, room_code=room_code,
                                   capacity=int(request.POST.get('capacity', 30)))
    messages.success(request, 'Đã tạo phòng thi thành công.')
    return JsonResponse({'success': True, 'room_id': room.id})


@login_required
@require_POST
def lecturer_delete_room(request, room_id):
    if not request.user.userprofile.is_lecturer: return JsonResponse({'success': False}, status=403)
    get_object_or_404(ExamRoom, pk=room_id).delete()
    messages.success(request, 'Đã xoá phòng thi thành công.')
    return JsonResponse({'success': True})


# Profile giảng viên
@login_required
def lecturer_profile_view(request: HttpRequest) -> HttpResponse:
    # Nếu không phải giảng viên, đẩy về trang profile của sinh viên
    if not request.user.userprofile.is_lecturer:
        return redirect('qna:profile')

    profile = request.user.userprofile
    avatar_url = _get_avatar_data_url(profile)

    context = {
        'user': request.user,
        'profile': profile,
        'avatar_url': avatar_url,
        'subjects_taught': profile.subjects_taught.all().order_by('subject_code')
    }
    return render(request, 'qna/lecturer/lecturer_profile.html', context)

@login_required
def lecturer_export_questions_word(request):
    if not request.user.userprofile.is_lecturer:
        return HttpResponse("Không có quyền.", status=403)

    subject_id = (request.GET.get("subject_id") or "").strip()
    difficulty = (request.GET.get("difficulty") or "ALL").strip().upper()

    if not subject_id or subject_id == "undefined":
        return JsonResponse(
            {"status": "FAIL", "message": "Thiếu hoặc sai subject_id."},
            status=400
        )

    try:
        subject_id = int(subject_id)
    except (TypeError, ValueError):
        return JsonResponse(
            {"status": "FAIL", "message": "subject_id không hợp lệ."},
            status=400
        )

    subject = get_object_or_404(
        request.user.userprofile.subjects_taught.all(),
        id=subject_id
    )

    questions = Question.objects.filter(subject=subject).exclude(
        question_id_in_barem__startswith="DRAFT_"
    )

    if difficulty in ["EASY", "MEDIUM", "HARD"]:
        questions = questions.filter(difficulty=difficulty)

    questions = questions.order_by("question_id_in_barem", "id")

    document = Document()

    # Thiết lập lề
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)

    # Style mặc định
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    # Header
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("TRƯỜNG ................................................")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("KHOA/BỘ MÔN ........................................")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ĐỀ THI VẤN ĐÁP")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    document.add_paragraph("")

    # Thông tin môn
    p = document.add_paragraph()
    run = p.add_run(f"Môn học: {subject.name}")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    p = document.add_paragraph()
    run = p.add_run(f"Mã môn: {subject.subject_code}")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    p = document.add_paragraph()
    run = p.add_run("Hình thức: Vấn đáp")
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    p = document.add_paragraph()
    run = p.add_run("Thời gian: ............ phút")
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    document.add_paragraph("")

    # Danh sách câu hỏi
    if not questions.exists():
        p = document.add_paragraph()
        run = p.add_run("Chưa có câu hỏi để xuất.")
        run.italic = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    else:
        for index, q in enumerate(questions, 1):
            p = document.add_paragraph()

            run1 = p.add_run(f"Câu {index}. ")
            run1.bold = True
            run1.font.name = "Times New Roman"
            run1.font.size = Pt(13)
            run1._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

            run2 = p.add_run(q.question_text)
            run2.font.name = "Times New Roman"
            run2.font.size = Pt(13)
            run2._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

            p2 = document.add_paragraph()
            run = p2.add_run("...............................................................................")
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

            p3 = document.add_paragraph()
            run = p3.add_run("...............................................................................")
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    document.add_paragraph("")
    document.add_paragraph("")

    # Chữ ký
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("GIẢNG VIÊN RA ĐỀ")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("(Ký và ghi rõ họ tên)")
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    filename = f"De_thi_{subject.subject_code}.docx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    document.save(response)
    return response